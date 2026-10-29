"""Separate inventory importer; dry-run by default. Never calls tra-onedrive-sync."""
import argparse
import csv
from datetime import datetime, date
from decimal import Decimal
import hashlib
import io
import json
import os
from pathlib import Path
import tempfile
import urllib.parse
import urllib.request

from prepare_inventory import FIELDS, json_exact, prepare

class NoRedirect(urllib.request.HTTPRedirectHandler):
    def redirect_request(self, req, fp, code, msg, headers, newurl):
        return None

# Do not forward OAuth/service credentials to redirect destinations.
OPENER = urllib.request.build_opener(NoRedirect)

def request(url, *, headers=None, data=None, limit=50_000_000):
    if urllib.parse.urlparse(url).scheme != 'https':
        raise ValueError('HTTPS required')
    req = urllib.request.Request(url, headers=headers or {}, data=data)
    with OPENER.open(req, timeout=60) as response:
        body = response.read(limit + 1)
        if len(body) > limit:
            raise ValueError('Response exceeds size limit')
        return body


def fetch_workbook(config):
    tenant = os.environ['TRA_INVENTORY_TENANT_ID']
    if not tenant or any(c not in 'abcdefghijklmnopqrstuvwxyzABCDEFGHIJKLMNOPQRSTUVWXYZ0123456789-.' for c in tenant):
        raise ValueError('Invalid tenant selector')
    token = json.loads(request('https://login.microsoftonline.com/' + tenant + '/oauth2/v2.0/token', headers={'Content-Type': 'application/x-www-form-urlencoded'}, data=urllib.parse.urlencode({'grant_type': 'client_credentials', 'client_id': os.environ['TRA_INVENTORY_CLIENT_ID'], 'client_secret': os.environ['TRA_INVENTORY_CLIENT_SECRET'], 'scope': 'https://graph.microsoft.com/.default'}).encode()))['access_token']
    metadata_url = config['graph_item_url']
    parsed = urllib.parse.urlparse(metadata_url)
    if parsed.scheme != 'https' or parsed.netloc != 'graph.microsoft.com' or not parsed.path.startswith('/v1.0/drives/'):
        raise ValueError('Expected a Microsoft Graph v1.0 drive item URL')
    metadata = json.loads(request(metadata_url, headers={'Authorization': 'Bearer ' + token}))
    download = metadata['@microsoft.graph.downloadUrl']
    host = urllib.parse.urlparse(download).hostname
    if host not in config['download_hosts']:
        raise ValueError('Workbook download host not approved in source configuration')
    # Download URL is short-lived authorization; never log or persist it.
    return request(download)


def workbook_to_csv(content, config, target):
    from openpyxl import load_workbook
    from openpyxl.utils.datetime import to_excel, WINDOWS_EPOCH
    columns = config['columns']
    if set(columns) != set(FIELDS):
        raise ValueError('Source configuration must map all 33 approved fields')
    indices = list(columns.values())
    if any(type(i) is not int or i < 1 for i in indices) or len(set(indices)) != 33:
        raise ValueError('Source column indices must be positive and distinct')
    if 16 in indices or 19 in indices:
        raise ValueError('Notes column P and Tenant Name column S are prohibited')
    verified = {'building': 1, 'unit_id': 2, 'floor': 3, 'unit_type': 4, 'bedrooms': 5, 'bathrooms': 6, 'square_feet': 7, 'program': 10, 'status': 11, 'gross_rent': 12, 'lease_start': 22, 'early_cancellation_option': 23, 'lease_expiration': 24, 'actual_occupancy': 25}
    if any(columns[k] != index for k, index in verified.items()):
        raise ValueError('Column mapping differs from verified production mapping')
    positions = {}
    workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    try:
        if workbook.epoch != WINDOWS_EPOCH:
            raise ValueError('Expected verified Excel 1900 date system')
        sheet = workbook['Master Inventory']
        header_row = config['header_row']
        if type(header_row) is not int or header_row < 1:
            raise ValueError('Invalid header row')
        header = next(sheet.iter_rows(min_row=header_row, max_row=header_row, values_only=True))
        if header[0] != 'Building' or header[1] != 'Unit ID':
            raise ValueError('Source header locator mismatch')
        if set(config['expected_headers']) != set(FIELDS):
            raise ValueError('Explicit source headers required for all fields')
        if any(header[index-1] != config['expected_headers'][field] for field, index in columns.items()):
            raise ValueError('Workbook source headers changed')
        with open(target, 'w', newline='', encoding='utf-8') as stream:
            writer = csv.DictWriter(stream, fieldnames=FIELDS)
            writer.writeheader()
            for row_number, values in enumerate(sheet.iter_rows(min_row=header_row+1, values_only=True), header_row+1):
                building, unit = values[0:2]
                if building is None and unit is None:
                    continue
                if building is None or unit is None:
                    raise ValueError('Incomplete workbook unit identity')
                positions[(str(building).strip(), str(unit).strip())] = row_number
                row = {}
                for field, index in columns.items():
                    value = values[index-1] if index <= len(values) else None
                    if isinstance(value, (date, datetime)):
                        if FIELDS[field][1] != 'date':
                            raise ValueError('Unexpected date in non-date field')
                        value = to_excel(value, epoch=WINDOWS_EPOCH)
                    row[field] = '' if value is None else value
                writer.writerow(row)
    finally:
        workbook.close()
    return positions


def publish(rows, report, expected):
    if report['records'] != expected['total'] or report['categories'] != {'Market': expected['market'], 'Affordable': expected['affordable']}:
        raise ValueError('Complete inventory/category counts do not match approved expectations')
    base = os.environ['TRA_INVENTORY_SUPABASE_URL'].rstrip('/')
    if base != 'https://bluavamwhplpyuzeppfc.supabase.co':
        raise ValueError('Unexpected Supabase project')
    key = os.environ['TRA_INVENTORY_SERVICE_ROLE_KEY']
    payload = {'p_units': rows, 'p_issues': report['date_correction_flags'], 'p_revision': report.get('workbook_sha256', report['source_sha256']), 'p_total': expected['total'], 'p_market': expected['market'], 'p_affordable': expected['affordable']}
    return json.loads(request(base + '/rest/v1/rpc/tra_publish_inventory', headers={'apikey': key, 'Authorization': 'Bearer ' + key, 'Content-Type': 'application/json'}, data=json_exact(payload).encode()))


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    source = parser.add_mutually_exclusive_group(required=True)
    source.add_argument('--csv')
    source.add_argument('--workbook')
    source.add_argument('--fetch', action='store_true')
    parser.add_argument('--config', required=True)
    parser.add_argument('--output', required=True)
    parser.add_argument('--publish', action='store_true', help='Explicit remote write; requires validated deployment and credentials')
    args = parser.parse_args()
    config = json.loads(Path(args.config).read_text())
    with tempfile.TemporaryDirectory() as temp:
        path = args.csv
        positions = None
        if not path:
            content = fetch_workbook(config) if args.fetch else Path(args.workbook).read_bytes()
            path = str(Path(temp) / 'inventory.csv')
            positions = workbook_to_csv(content, config, path)
        rows, report = prepare(path)
        if positions is not None:
            from openpyxl.utils import get_column_letter
            source_headers = {target: source for source, (target, kind) in FIELDS.items()}
            for issue in report['date_correction_flags']:
                issue['workbook_cell'] = get_column_letter(config['columns'][source_headers[issue['field']]]) + str(positions[(issue['building'], issue['unit_id'])])
            report['workbook_sha256'] = hashlib.sha256(content).hexdigest()
        expected = config['expected_counts']
        if report['records'] != expected['total'] or report['categories'] != {'Market': expected['market'], 'Affordable': expected['affordable']}:
            raise ValueError('Inventory completeness check failed')
        output = Path(args.output)
        output.mkdir(parents=True, exist_ok=True)
        (output / 'inventory-candidate.json').write_text(json_exact(rows) + '\n')
        (output / 'inventory-validation.json').write_text(json.dumps(report, indent=2) + '\n')
        if args.publish:
            publish(rows, report, expected)
            print('Inventory batch published atomically')
        else:
            print('Dry run passed; no remote writes')

if __name__ == '__main__':
    try:
        main()
    except Exception as error:
        # HTTP errors may include signed URLs or authentication details. Never print them.
        if isinstance(error, ValueError):
            raise SystemExit(str(error)) from None
        raise SystemExit('Inventory import failed (' + type(error).__name__ + '); no activation performed') from None
