import io
import json
from pathlib import Path
import tempfile
import unittest
from openpyxl import Workbook
from openpyxl.utils.datetime import CALENDAR_MAC_1904
from prepare_inventory import FIELDS, prepare
from sync_inventory import workbook_to_csv

class WorkbookTests(unittest.TestCase):
    def fixture(self):
        verified = {'building':1,'unit_id':2,'floor':3,'unit_type':4,'bedrooms':5,'bathrooms':6,'square_feet':7,'program':10,'status':11,'gross_rent':12,'lease_start':22,'early_cancellation_option':23,'lease_expiration':24,'actual_occupancy':25}
        unused=(i for i in range(1,80) if i not in set(verified.values())|{16,19})
        columns={k:verified[k] if k in verified else next(unused) for k in FIELDS}
        wb=Workbook();ws=wb.active;ws.title='Master Inventory'
        for k,index in columns.items():
            ws.cell(1,index,{'building':'Building','unit_id':'Unit ID'}.get(k,k))
            value='' if FIELDS[k][1]=='text' else 0
            if FIELDS[k][1]=='date':value=50000
            ws.cell(2,index,value)
        for k,value in {'building':'TEST-BUILDING','unit_id':'TEST-0001','program':'Market','status':'Signed','lease_start':'invalid-date','actual_occupancy':16}.items():ws.cell(2,columns[k],value)
        # Sensitive source values must never reach the mapped CSV.
        ws.cell(1,16,'Notes');ws.cell(2,16,'PRIVATE_SENTINEL')
        ws.cell(1,19,'Tenant Name');ws.cell(2,19,'PRIVATE_SENTINEL')
        config={'columns':columns,'header_row':1,'expected_headers':{k:ws.cell(1,i).value for k,i in columns.items()}}
        return wb,config
    def test_workbook_conversion_retains_invalid_date_unit_and_excludes_private_columns(self):
        wb,config=self.fixture();stream=io.BytesIO();wb.save(stream)
        with tempfile.TemporaryDirectory() as tmp:
            path=Path(tmp)/'out.csv';workbook_to_csv(stream.getvalue(),config,path)
            self.assertNotIn('PRIVATE_SENTINEL',path.read_text())
            rows,report=prepare(path)
            self.assertEqual(len(rows),1)
            self.assertEqual(rows[0]['actual_occupancy'],16)
            self.assertEqual(rows[0]['early_cancellation_date'],'2036-11-21')
            self.assertIsNone(rows[0]['lease_start'])
            self.assertEqual(len(report['date_correction_flags']),1)
    def test_changed_header_and_epoch_rejected(self):
        for change in ['header','epoch']:
            wb,config=self.fixture()
            if change=='header':wb.active.cell(1,1,'Wrong header')
            else:wb.epoch=CALENDAR_MAC_1904
            stream=io.BytesIO();wb.save(stream)
            with tempfile.TemporaryDirectory() as tmp:
                with self.assertRaises(ValueError):workbook_to_csv(stream.getvalue(),config,Path(tmp)/'out.csv')

if __name__=='__main__':unittest.main()
